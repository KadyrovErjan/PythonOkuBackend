"""
Export view для аналитики
"""
from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from django.http import HttpResponse
import io
import logging

from ..permissions import CanViewAnalytics
from ..services.analytics import get_summary_analytics, get_monthly_analytics

logger = logging.getLogger(__name__)


@api_view(['GET'])
@permission_classes([IsAuthenticated, CanViewAnalytics])
def export_analytics(request):
    """
    Export analytics to Excel
    GET /api/v1/analytics/export/?format=xlsx
    """
    logger.info("=== EXPORT VIEW CALLED ===")
    logger.info(f"User: {request.user}, Method: {request.method}")
    
    from django.db import connection
    connection.close()
    
    try:
        from openpyxl import Workbook
        
        summary = get_summary_analytics()
        monthly = get_monthly_analytics()
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Аналитика"
        
        ws['A1'] = 'Сводка'
        ws['A2'] = 'Общий доход'
        ws['B2'] = str(summary['total_income'])
        ws['A3'] = 'Общий расход'
        ws['B3'] = str(summary['total_expense'])
        ws['A4'] = 'Чистая прибыль'
        ws['B4'] = str(summary['net_profit'])
        
        ws['A6'] = 'Месяц'
        ws['B6'] = 'Доход'
        ws['C6'] = 'Расход'
        ws['D6'] = 'Прибыль'
        
        row = 7
        for item in monthly:
            ws[f'A{row}'] = item['month']
            ws[f'B{row}'] = str(item['income'])
            ws[f'C{row}'] = str(item['expense'])
            ws[f'D{row}'] = str(item['profit'])
            row += 1
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        logger.info(f"Excel file created: {len(output.getvalue())} bytes")
        
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="analytics.xlsx"'
        logger.info("Returning Excel response")
        return response
        
    except Exception as e:
        logger.error(f"Export error: {e}", exc_info=True)
        return Response(
            {'detail': f'Ошибка экспорта: {str(e)}'},
            status=status.HTTP_500_INTERNAL_SERVER_ERROR
        )
